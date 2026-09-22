"""M6: Student roster import/export backend.

Covers EXPORT, PREVIEW, APPLY, M5 (ADR 0045 presentation-only) regression,
and permission-registry live-consumer regression scenarios described in the
M6 delegated task. Uses an isolated in-memory SQLite database and a minimal
FastAPI app mounting only ``routers/students.py``, mirroring the pattern used
by ``tests/test_students_ui.py`` and ``tests/test_talent_org_intelligence_queries.py``.
"""
from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import auth
import models
import student_roster_service as roster
from academic_grade import AL_ANDALUS_SECTION_DISPLAY_WORKSPACE_UUID
from auth import get_current_user
from database import Base
from dependencies import get_db
from routers import students as students_router
from student_academic_service import create_student_with_number, current_student_number


ROSTER_HEADERS = (
    "student_id", "first_name", "father_name", "last_name", "gender", "status",
    "branch", "academic_year", "grade", "section",
)


@pytest.fixture()
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    session.add_all((
        models.SchoolGroup(id=1, name="One", workspace_uuid="11111111-1111-1111-1111-111111111111"),
        models.SchoolGroup(id=2, name="Two", workspace_uuid="22222222-2222-2222-2222-222222222222"),
        models.SchoolGroup(id=3, name="AlAndalus", workspace_uuid=AL_ANDALUS_SECTION_DISPLAY_WORKSPACE_UUID),
    ))
    session.add_all((
        models.Branch(id=10, school_group_id=1, name="North", status=True),
        models.Branch(id=11, school_group_id=1, name="South", status=True),
        models.Branch(id=20, school_group_id=2, name="Foreign", status=True),
        models.Branch(id=30, school_group_id=3, name="Main", status=True),
    ))
    session.add_all((
        models.AcademicYear(id=100, school_group_id=1, year_name="2026-2027", is_active=True),
        models.AcademicYear(id=200, school_group_id=2, year_name="2026-2027", is_active=True),
        models.AcademicYear(id=300, school_group_id=3, year_name="2026-2027", is_active=True),
    ))
    session.add_all((
        models.PlanningSection(id=1, grade_level="1", section_name="A", class_status="active", branch_id=10, academic_year_id=100),
        models.PlanningSection(id=2, grade_level="1", section_name="B", class_status="active", branch_id=11, academic_year_id=100),
        models.PlanningSection(id=3, grade_level="1", section_name="A", class_status="active", branch_id=30, academic_year_id=300),
    ))
    session.commit()
    yield session
    session.close()


def actor(*, scope="ORGANIZATION", branch=10, role="Editor", group=1, uid="u1"):
    return models.User(
        user_id=uid, username=uid, role=role, user_type="TENANT", access_scope=scope,
        school_group_id=group, branch_id=branch, academic_year_id=100 if group == 1 else (200 if group == 2 else 300),
        is_active=True,
    )


def permissions(db, *keys, role="Editor", group=1, allow=True):
    for key in dict.fromkeys(keys):
        row = db.query(models.RolePermission).filter_by(
            school_group_id=group, role=role, permission_key=key,
        ).one_or_none()
        if row is None:
            db.add(models.RolePermission(school_group_id=group, role=role, permission_key=key, is_allowed=allow))
        else:
            row.is_allowed = allow
    db.commit()


@pytest.fixture()
def client(db):
    app = FastAPI()
    app.include_router(students_router.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: actor()
    return TestClient(app)


def _workbook_bytes(rows, headers=ROSTER_HEADERS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(headers))
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.read()


def _row(student_id="0000000001", first="Ada", father="", last="Lovelace", gender="Female",
         status="active", branch="North", year="2026-2027", grade="1", section="A"):
    return [student_id, first, father, last, gender, status, branch, year, grade, section]


def _upload(client, path, content, filename="roster.xlsx"):
    return client.post(path, files={"roster_file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


# ---------------------------------------------------------------------------
# EXPORT
# ---------------------------------------------------------------------------

def test_export_requires_permission(client):
    response = client.get("/api/students/roster/export")
    assert response.status_code == 403


def test_export_authorized_returns_xlsx_with_deterministic_columns(db, client):
    permissions(db, "students.export", "students.view")
    create_student_with_number(db, school_group_id=1, student_number="0000000001", first_name="Ada", last_name="Lovelace", actor=actor())
    db.commit()
    response = client.get("/api/students/roster/export")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert headers == list(roster.EXPORT_HEADERS)


def test_export_preserves_student_id_leading_zeros_as_text(db, client):
    permissions(db, "students.export", "students.view")
    create_student_with_number(db, school_group_id=1, student_number="0000000123", first_name="Ada", last_name="Lovelace", actor=actor())
    db.commit()
    response = client.get("/api/students/roster/export")
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    cell = sheet.cell(row=2, column=1)
    assert cell.value == "STD0000000123"
    assert isinstance(cell.value, str)
    assert cell.number_format == "@"


def test_export_legacy_student_with_no_managed_id_has_blank_id_cell(db, client):
    permissions(db, "students.export", "students.view")
    db.add(models.Student(school_group_id=1, first_name="Legacy", last_name="Student", status="active"))
    db.commit()
    response = client.get("/api/students/roster/export")
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    row = [cell.value for cell in sheet[2]]
    assert row[0] in (None, "")
    assert row[1] == "Legacy"


def test_export_is_tenant_isolated(db, client):
    permissions(db, "students.export", "students.view")
    create_student_with_number(db, school_group_id=1, student_number="0000000001", first_name="Mine", last_name="One", actor=actor())
    create_student_with_number(db, school_group_id=2, student_number="0000000002", first_name="Other", last_name="Two", actor=actor(group=2))
    db.commit()
    response = client.get("/api/students/roster/export")
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    names = [cell.value for row in sheet.iter_rows(min_row=2) for cell in [row[1]]]
    assert "Mine" in names
    assert "Other" not in names


def test_export_branch_scoped_actor_only_sees_own_branch_students(db, client):
    permissions(db, "students.export", "students.view")
    north = create_student_with_number(db, school_group_id=1, student_number="0000000010", first_name="North", last_name="Kid", actor=actor())
    south = create_student_with_number(db, school_group_id=1, student_number="0000000011", first_name="South", last_name="Kid", actor=actor())
    db.commit()
    from student_academic_service import create_placement
    create_placement(db, school_group_id=1, student_id=north.id, academic_year_id=100, branch_id=10,
                      planning_section_id=1, grade_level="1", section_name="A", effective_from=datetime(2026, 9, 1))
    create_placement(db, school_group_id=1, student_id=south.id, academic_year_id=100, branch_id=11,
                      planning_section_id=2, grade_level="1", section_name="B", effective_from=datetime(2026, 9, 1))
    db.commit()

    client.app.dependency_overrides[get_current_user] = lambda: actor(scope="BRANCH", branch=10)
    response = client.get("/api/students/roster/export")
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    names = [cell.value for row in sheet.iter_rows(min_row=2) for cell in [row[1]]]
    assert "North" in names
    assert "South" not in names


# ---------------------------------------------------------------------------
# PREVIEW
# ---------------------------------------------------------------------------

def test_preview_requires_permission(client):
    response = _upload(client, "/api/students/roster/import/preview", _workbook_bytes([_row()]))
    assert response.status_code == 403


def test_preview_valid_workbook_returns_ok_rows_with_zero_mutation(db, client):
    permissions(db, "students.import")
    before = db.query(models.Student).count()
    response = _upload(client, "/api/students/roster/import/preview", _workbook_bytes([_row()]))
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ok"
    assert body["summary"] == {"total_rows": 1, "valid_rows": 1, "error_rows": 0}
    assert db.query(models.Student).count() == before


def test_preview_rejects_non_xlsx_extension(db, client):
    permissions(db, "students.import")
    response = _upload(client, "/api/students/roster/import/preview", b"not a workbook", filename="roster.csv")
    assert response.status_code == 422
    assert response.json()["file_error"]["error_code"] == "invalid_extension"


def test_preview_rejects_malformed_workbook(db, client):
    permissions(db, "students.import")
    response = _upload(client, "/api/students/roster/import/preview", b"this is not an xlsx file at all")
    assert response.status_code == 422
    assert response.json()["file_error"]["error_code"] == "workbook_unreadable"


def test_preview_rejects_missing_required_column(db, client):
    permissions(db, "students.import")
    headers = tuple(h for h in ROSTER_HEADERS if h != "branch")
    content = _workbook_bytes([[c for c in _row() if c != "North"]], headers=headers)
    response = _upload(client, "/api/students/roster/import/preview", content)
    assert response.status_code == 422
    assert response.json()["file_error"]["error_code"] == "missing_required_column"


def test_preview_rejects_unexpected_column(db, client):
    permissions(db, "students.import")
    headers = ROSTER_HEADERS + ("unexpected_column",)
    content = _workbook_bytes([_row() + ["x"]], headers=headers)
    response = _upload(client, "/api/students/roster/import/preview", content)
    assert response.status_code == 422
    assert response.json()["file_error"]["error_code"] == "unexpected_column"


def test_preview_flags_duplicate_student_id_within_workbook(db, client):
    permissions(db, "students.import")
    content = _workbook_bytes([_row(student_id="0000000001"), _row(student_id="0000000001", first="Second")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    body = response.json()
    assert body["summary"]["error_rows"] == 1
    second_row = body["rows"][1]
    assert second_row["status"] == "error"
    assert any(e["error_code"] == "duplicate_student_id_in_file" for e in second_row["errors"])


def test_preview_flags_invalid_student_id_format(db, client):
    permissions(db, "students.import")
    content = _workbook_bytes([_row(student_id="ABC")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    errors = response.json()["rows"][0]["errors"]
    assert any(e["error_code"] == "invalid_student_id" for e in errors)


def test_preview_accepts_student_id_leading_zeros(db, client):
    permissions(db, "students.import")
    content = _workbook_bytes([_row(student_id="0000000009")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    assert response.json()["rows"][0]["status"] == "ok"
    assert response.json()["rows"][0]["data"]["student_number"] == "0000000009"


def test_preview_same_tenant_conflict_reveals_minimal_identity(db, client):
    permissions(db, "students.import", "students.view")
    create_student_with_number(db, school_group_id=1, student_number="0000000001", first_name="Existing", last_name="Person", actor=actor())
    db.commit()
    content = _workbook_bytes([_row(student_id="0000000001")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    errors = response.json()["rows"][0]["errors"]
    conflict = next(e for e in errors if e["error_code"] == "student_id_conflict")
    assert conflict["display_name"] == "Existing Person"
    assert "student_id" in conflict


def test_preview_same_tenant_conflict_without_view_permission_is_generic(db, client):
    permissions(db, "students.import")
    create_student_with_number(db, school_group_id=1, student_number="0000000001", first_name="Existing", last_name="Person", actor=actor())
    db.commit()
    content = _workbook_bytes([_row(student_id="0000000001")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    errors = response.json()["rows"][0]["errors"]
    conflict = next(e for e in errors if e["error_code"] == "student_id_conflict")
    assert "student_id" not in conflict
    assert "display_name" not in conflict


def test_preview_cross_tenant_conflict_is_fully_generic(db, client):
    permissions(db, "students.import", "students.view")
    create_student_with_number(db, school_group_id=2, student_number="0000000001", first_name="Foreign", last_name="Person", actor=actor(group=2))
    db.commit()
    content = _workbook_bytes([_row(student_id="0000000001")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    errors = response.json()["rows"][0]["errors"]
    conflict = next(e for e in errors if e["error_code"] == "student_id_conflict")
    assert "student_id" not in conflict
    assert "display_name" not in conflict
    assert "Foreign" not in str(response.json())


def test_preview_flags_invalid_branch(db, client):
    permissions(db, "students.import")
    content = _workbook_bytes([_row(branch="Nonexistent")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    errors = response.json()["rows"][0]["errors"]
    assert any(e["error_code"] == "invalid_branch" for e in errors)


def test_preview_flags_cross_tenant_branch_reference_as_invalid_not_leaked(db, client):
    permissions(db, "students.import")
    content = _workbook_bytes([_row(branch="Foreign")])  # belongs to school_group_id=2
    response = _upload(client, "/api/students/roster/import/preview", content)
    errors = response.json()["rows"][0]["errors"]
    assert any(e["error_code"] == "invalid_branch" for e in errors)


def test_preview_flags_branch_outside_actor_scope_as_unauthorized(db, client):
    permissions(db, "students.import")
    client.app.dependency_overrides[get_current_user] = lambda: actor(scope="BRANCH", branch=10)
    content = _workbook_bytes([_row(branch="South")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    errors = response.json()["rows"][0]["errors"]
    assert any(e["error_code"] == "unauthorized_reference" for e in errors)


def test_preview_flags_invalid_academic_year(db, client):
    permissions(db, "students.import")
    content = _workbook_bytes([_row(year="1999-2000")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    errors = response.json()["rows"][0]["errors"]
    assert any(e["error_code"] == "invalid_academic_year" for e in errors)


def test_preview_flags_invalid_grade(db, client):
    permissions(db, "students.import")
    content = _workbook_bytes([_row(grade="99")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    errors = response.json()["rows"][0]["errors"]
    assert any(e["error_code"] == "invalid_grade" for e in errors)


def test_preview_flags_invalid_section(db, client):
    permissions(db, "students.import")
    content = _workbook_bytes([_row(section="Z")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    errors = response.json()["rows"][0]["errors"]
    assert any(e["error_code"] == "invalid_section" for e in errors)


def test_preview_flags_missing_values(db, client):
    permissions(db, "students.import")
    content = _workbook_bytes([_row(first="")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    errors = response.json()["rows"][0]["errors"]
    assert any(e["error_code"] == "missing_value" and e["field"] == "first_name" for e in errors)


def test_preview_flags_invalid_gender_value(db, client):
    permissions(db, "students.import")
    content = _workbook_bytes([_row(gender="Other")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    errors = response.json()["rows"][0]["errors"]
    assert any(e["error_code"] == "invalid_field" and e["field"] == "gender" for e in errors)


def test_preview_handles_formula_cell_safely_without_evaluation(db, client):
    permissions(db, "students.import")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(ROSTER_HEADERS))
    sheet.append(_row())
    sheet.cell(row=2, column=1).value = "=1+1"  # never opened in Excel: no cached value
    output = BytesIO(); workbook.save(output); output.seek(0)
    response = _upload(client, "/api/students/roster/import/preview", output.read())
    assert response.status_code in (200, 422)
    body = response.json()
    errors = body["rows"][0].get("errors", [])
    assert body["rows"][0]["status"] == "error"
    assert any(e["error_code"] == "missing_value" for e in errors)


def test_preview_rejects_oversized_row_count(db, client, monkeypatch):
    permissions(db, "students.import")
    monkeypatch.setattr(roster, "MAX_DATA_ROWS", 2)
    content = _workbook_bytes([_row(student_id=f"{i:010d}") for i in range(5)])
    response = _upload(client, "/api/students/roster/import/preview", content)
    assert response.status_code == 422
    assert response.json()["file_error"]["error_code"] == "too_many_rows"


# ---------------------------------------------------------------------------
# APPLY
# ---------------------------------------------------------------------------

def test_apply_requires_permission(client):
    response = _upload(client, "/api/students/roster/import/apply", _workbook_bytes([_row()]))
    assert response.status_code == 403


def test_apply_valid_workbook_creates_student_and_placement_atomically(db, client):
    permissions(db, "students.import")
    response = _upload(client, "/api/students/roster/import/apply", _workbook_bytes([_row(student_id="0000000042")]))
    assert response.status_code == 201
    body = response.json()
    assert body["applied"] is True
    assert len(body["created_student_ids"]) == 1
    student_id = body["created_student_ids"][0]
    assert current_student_number(db, school_group_id=1, student_id=student_id) == "STD0000000042"
    placement = db.query(models.StudentAcademicPlacement).filter_by(student_id=student_id).one()
    assert (placement.branch_id, placement.grade_level, placement.section_name) == (10, "1", "A")


def test_apply_does_not_trust_preview_and_revalidates_independently(db, client):
    permissions(db, "students.import")
    # No preview call at all - apply must independently parse/validate.
    response = _upload(client, "/api/students/roster/import/apply", _workbook_bytes([_row(student_id="0000000043")]))
    assert response.status_code == 201


def test_apply_rolls_back_entire_batch_when_one_row_is_invalid(db, client):
    permissions(db, "students.import")
    before = db.query(models.Student).count()
    content = _workbook_bytes([_row(student_id="0000000050"), _row(student_id="0000000051", branch="Nonexistent")])
    response = _upload(client, "/api/students/roster/import/apply", content)
    assert response.status_code == 422
    body = response.json()
    assert body["applied"] is False
    assert db.query(models.Student).count() == before


def test_apply_detects_student_id_race_after_preview_and_rolls_back(db, client):
    permissions(db, "students.import")
    content = _workbook_bytes([_row(student_id="0000000060")])
    preview = _upload(client, "/api/students/roster/import/preview", content)
    assert preview.json()["status"] == "ok"
    # Someone else claims the same Student ID between preview and apply.
    create_student_with_number(db, school_group_id=1, student_number="0000000060", first_name="Race", last_name="Winner", actor=actor())
    db.commit()
    before = db.query(models.Student).count()
    apply_response = _upload(client, "/api/students/roster/import/apply", content)
    assert apply_response.status_code == 422
    assert apply_response.json()["applied"] is False
    assert db.query(models.Student).count() == before


def test_apply_rechecks_permission_and_scope(db, client):
    # No students.import permission granted at all.
    content = _workbook_bytes([_row(student_id="0000000070")])
    response = _upload(client, "/api/students/roster/import/apply", content)
    assert response.status_code == 403


def test_apply_rechecks_branch_scope(db, client):
    permissions(db, "students.import")
    client.app.dependency_overrides[get_current_user] = lambda: actor(scope="BRANCH", branch=10)
    content = _workbook_bytes([_row(student_id="0000000071", branch="South")])
    response = _upload(client, "/api/students/roster/import/apply", content)
    assert response.status_code == 422
    errors = response.json()["rows"][0]["errors"]
    assert any(e["error_code"] == "unauthorized_reference" for e in errors)


def test_apply_does_not_corrupt_unrelated_historical_placement(db, client):
    permissions(db, "students.import")
    other = create_student_with_number(db, school_group_id=1, student_number="0000000080", first_name="Other", last_name="Existing", actor=actor())
    db.commit()
    from student_academic_service import create_placement
    create_placement(db, school_group_id=1, student_id=other.id, academic_year_id=100, branch_id=10,
                      planning_section_id=1, grade_level="1", section_name="A", effective_from=datetime(2026, 1, 1))
    db.commit()
    before_placement = db.query(models.StudentAcademicPlacement).filter_by(student_id=other.id).one()
    before_effective_from = before_placement.effective_from

    content = _workbook_bytes([_row(student_id="0000000081")])
    response = _upload(client, "/api/students/roster/import/apply", content)
    assert response.status_code == 201

    after_placement = db.query(models.StudentAcademicPlacement).filter_by(student_id=other.id).one()
    assert after_placement.effective_from == before_effective_from


# ---------------------------------------------------------------------------
# M5 / ADR 0045 regression
# ---------------------------------------------------------------------------

def test_import_never_matches_section_by_display_label(db, client):
    permissions(db, "students.import", group=3)
    client.app.dependency_overrides[get_current_user] = lambda: actor(group=3, branch=30)
    content = _workbook_bytes([_row(student_id="0000000090", branch="Main", section="1.1")])
    response = _upload(client, "/api/students/roster/import/preview", content)
    errors = response.json()["rows"][0]["errors"]
    assert any(e["error_code"] == "invalid_section" for e in errors)


def test_export_section_display_is_additive_only_on_al_andalus_workspace(db, client):
    permissions(db, "students.export", "students.view", group=3)
    client.app.dependency_overrides[get_current_user] = lambda: actor(group=3, branch=30)
    student = create_student_with_number(db, school_group_id=3, student_number="0000000091", first_name="Layla", last_name="Student", actor=actor(group=3))
    db.commit()
    from student_academic_service import create_placement
    create_placement(db, school_group_id=3, student_id=student.id, academic_year_id=300, branch_id=30,
                      planning_section_id=3, grade_level="1", section_name="A", effective_from=datetime(2026, 9, 1))
    db.commit()
    response = client.get("/api/students/roster/export")
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    row = [cell.value for cell in sheet[2]]
    section_idx = roster.EXPORT_HEADERS.index("section")
    display_idx = roster.EXPORT_HEADERS.index("section_display")
    assert row[section_idx] == "A"
    assert row[display_idx] == "1.1"


def test_export_other_tenant_section_display_falls_back_to_canonical(db, client):
    permissions(db, "students.export", "students.view")
    student = create_student_with_number(db, school_group_id=1, student_number="0000000092", first_name="Kim", last_name="Student", actor=actor())
    db.commit()
    from student_academic_service import create_placement
    create_placement(db, school_group_id=1, student_id=student.id, academic_year_id=100, branch_id=10,
                      planning_section_id=1, grade_level="1", section_name="A", effective_from=datetime(2026, 9, 1))
    db.commit()
    response = client.get("/api/students/roster/export")
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    row = [cell.value for cell in sheet[2]]
    section_idx = roster.EXPORT_HEADERS.index("section")
    display_idx = roster.EXPORT_HEADERS.index("section_display")
    assert row[section_idx] == "A"
    assert row[display_idx] == "A"


# ---------------------------------------------------------------------------
# Permission registry regression
# ---------------------------------------------------------------------------

def test_students_import_export_permissions_are_registered():
    import permission_registry as pr
    assert "students.import" in pr.ALL_PERMISSION_KEYS
    assert "students.export" in pr.ALL_PERMISSION_KEYS
    assert "students.import" in pr.PERMISSION_LABELS
    assert "students.export" in pr.PERMISSION_LABELS
